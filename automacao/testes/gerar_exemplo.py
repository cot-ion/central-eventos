"""
Gera um .pptx sintético no padrão do infográfico (páginas 1 e 16 + Anotações),
usado pelos testes do extrator. Coordenadas em "pixels" de um slide 1330x750.

    python testes/gerar_exemplo.py saida.pptx
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ESCALA = 12192000 / 1330  # px → EMU (slide 16:9)
CINZA = RGBColor(0xD9, 0xD9, 0xD9)
AZUL = RGBColor(0x5B, 0x9B, 0xD5)
VERMELHO = RGBColor(0xFF, 0x00, 0x00)


def e(v):
    return Emu(int(v * ESCALA))


def texto(shapes, x, y, w, h, txt, tam=9, cor=None, riscado=False):
    tb = shapes.add_textbox(e(x), e(y), e(w), e(h))
    tf = tb.text_frame
    for i, linha in enumerate(txt.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = linha
        r.font.size = Pt(tam)
        if cor:
            r.font.color.rgb = cor
        if riscado:
            r.font._rPr.set("strike", "sngStrike")
    return tb


def caixa(shapes, x, y, w, h, txt="", fill=None, cor=None, riscado=False, forma=MSO_SHAPE.RECTANGLE):
    sh = shapes.add_shape(forma, e(x), e(y), e(w), e(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if txt:
        tf = sh.text_frame
        for i, linha in enumerate(txt.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = linha
            r.font.size = Pt(8)
            if cor:
                r.font.color.rgb = cor
            if riscado:
                r.font._rPr.set("strike", "sngStrike")
    return sh


def seta(shapes, x1, y1, x2, y2):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, e(x1), e(y1), e(x2), e(y2))
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    ln.append(tail)
    return c


def no(shapes, x, y, txt, cor=None, alto=225):
    return caixa(shapes, x, y, 100, alto, txt, cor=cor, forma=MSO_SHAPE.ROUNDED_RECTANGLE)


def sinal(shapes, x, y, nome, via=None, direcao=1, ate=None, riscado=False, cor=None):
    """Pílula em (x, y) com seta para a direita (1) ou esquerda (-1) até `ate`."""
    caixa(shapes, x, y, 72, 22, nome, fill=CINZA, riscado=riscado, cor=cor)
    cy = y + 11
    if direcao > 0:
        seta(shapes, x + 72, cy, ate, cy)
        if via:
            texto(shapes, x + 80, cy - 16, 160, 14, via, tam=7)
    else:
        seta(shapes, x, cy, ate, cy)
        if via:
            texto(shapes, ate + 100, cy - 16, 160, 14, via, tam=7)


def slide_ion_sp(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    texto(s, 15, 10, 300, 40, "IONxSP", tam=24)
    texto(s, 5, 72, 150, 75, "Modelo:\nEstúdio: 6 SP\nControle: ION", tam=9)
    caixa(s, 552, 107, 100, 28, "VIDEO", fill=AZUL)
    caixa(s, 552, 414, 100, 28, "AUDIO", fill=AZUL)

    # ---- vídeo
    no(s, 88, 163, "COT\nSP")
    no(s, 502, 163, "COT ION")
    no(s, 838, 163, "CTL ION")
    no(s, 1122, 163, "EXIBE")
    for i, (cam, via) in enumerate([("CAMERA 1", "PR RX SP-1"), ("CAMERA 2", "PR RX SP-2"), ("CAMERA 3", "PR RX SP-3")]):
        sinal(s, 195, 165 + 34 * i, cam, via, 1, 497)
    for i, nome in enumerate(["DIRETO TIT", "AUX TELÃO", "PGM COM\nCTL"]):
        sinal(s, 421, 286 + 33 * i, nome, "PR TX LAWO E DR LAWO", -1, 188)
    # coluna 2 dentro de um grupo (testa coordenadas de grupo)
    g = s.add_group_shape()
    for i, nome in enumerate(["CAMERA 1", "CAMERA 2", "CAMERA 3", "DIRETO TIT", "DIRETO BY"]):
        sinal(g.shapes, 611, 165 + 34 * i, nome, None, 1, 833)
    sinal(s, 760, 354, "AUX TELÃO", None, -1, 603)
    for i, nome in enumerate(["PGM COM", "PGM PROD"]):
        sinal(s, 942, 164 + 34 * i, nome, None, 1, 1118)
    caixa(s, 942, 231, 72, 22, "PGM\nLIMPO", fill=CINZA)
    seta(s, 978, 253, 978, 292)
    texto(s, 955, 292, 50, 14, "GRAVAR", tam=7)

    # ---- áudio
    no(s, 88, 460, "COT\nSP")
    no(s, 500, 460, "COT ION")
    no(s, 840, 460, "CTL ION")
    for i, (nome, via) in enumerate([("MIC BY", "DIRECT"), ("OVER", "DIRECT"), ("+ASS SP", "PARTYLINE"), ("+CAM SP", "PARTYLINE")]):
        sinal(s, 195, 463 + 33 * i, nome, via, 1, 497)
    sinal(s, 425, 648, "MINUS", "DIRECT", -1, 188)
    for i, (nome, via) in enumerate([("MIC BY", "REMOTE"), ("OVER", "IFB EVNT"), ("+CAM SP", None), ("+ASS SP", None)]):
        sinal(s, 613, 465 + 34 * i, nome, via, 1, 836)
    sinal(s, 763, 648, "MINUS", "AUX AUD", -1, 602)

    # ---- canais
    texto(s, 985, 425, 160, 90, "Câmeras 1,2 e 3 SP:\nCH 1 – LAPELA NARRA\nCH 2 – LAPELA COMENTA\nCH 3 – MD NARRA\nCH 4 – MD COMENTA", tam=7)
    texto(s, 1160, 425, 160, 80, "PGM Comercializado:\n" + "\n".join(f"CH {i} – MIX COM FOGUETE" for i in range(1, 5)), tam=7)
    texto(s, 1160, 515, 160, 80, "PGM Produzido:\n" + "\n".join(f"CH {i} – MIX SEM FOGUETE" for i in range(1, 5)), tam=7, cor=VERMELHO)
    sl = prs.slides[-1]
    sl.notes_slide.notes_text_frame.text = "Nome: SP · Estúdio 6 SP com controle ION\nApelidos: câmeras sp, narração sp"


def slide_tvg(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    texto(s, 15, 10, 300, 40, "ION x SITE TVG", tam=24)
    texto(s, 12, 72, 150, 75, "Modelo: TVG\nEstúdio: JB/SP/RCE/BHE\nControle: JB/SP/RCE/BHE", tam=9)
    caixa(s, 552, 107, 100, 28, "VIDEO", fill=AZUL)
    caixa(s, 552, 414, 100, 28, "AUDIO", fill=AZUL)
    no(s, 88, 163, "COT\nEXT")
    no(s, 502, 163, "COT ION")
    texto(s, 815, 170, 25, 110, "G\nR\nA\nV\nA\nR", tam=9)
    sinal(s, 195, 225, "PGM PROD", None, 1, 497)
    sinal(s, 195, 258, "PGM LIMPO\nSEM SONO", None, 1, 497)
    sinal(s, 421, 352, "DIRETO", "LAWO", -1, 188)
    sinal(s, 611, 175, "DIRETO", None, 1, 788)
    sinal(s, 611, 209, "PGM PROD", None, 1, 788)
    sinal(s, 611, 243, "PGM LIMPO\nSEM SONO", None, 1, 788, riscado=True)
    no(s, 88, 460, "COT\nEXT")
    no(s, 500, 460, "COT ION")
    sinal(s, 195, 463, "CORD", "DIRECT", 1, 497)
    sinal(s, 425, 653, "MINUS", "DIRECT", -1, 188)
    texto(s, 130, 698, 400, 14, "*Talkback (ch1), Foldback (ch2), Ponto 1 (ch3) e Ponto 2 (ch4)", tam=7)


def slide_anotacoes(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    texto(s, 15, 10, 300, 40, "Anotações", tam=24)
    texto(s, 140, 140, 500, 120, "Estudio j – backstage/ eliminação\nCtl MG3 - Porchat ao vivo", tam=14)


def slide_capa(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    texto(s, 400, 300, 500, 100, "globo", tam=60)


def gerar(destino: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    slide_ion_sp(prs)
    slide_tvg(prs)
    slide_anotacoes(prs)
    slide_capa(prs)
    prs.save(str(destino))
    return destino


if __name__ == "__main__":
    print(gerar(Path(sys.argv[1] if len(sys.argv) > 1 else "exemplo.pptx")))
