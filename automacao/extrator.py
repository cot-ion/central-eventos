"""
Extrai os modelos de evento de um .pptx da Central de Eventos.

Cada slide no padrão do infográfico vira um evento:

    ┌─────────┐  CAMERA 1 ─PR RX SP-1──▶  ┌─────────┐ ...
    │ COT SP  │                           │ COT ION │
    └─────────┘  ◀──PR TX LAWO── DIRETO   └─────────┘

A leitura é geométrica, por isso não depende da ordem em que as formas
foram criadas no PowerPoint:
  * blocos ("nós")  = caixas de texto altas (COT SP, CTL ION, EXIBE…);
  * sinais          = caixas pequenas preenchidas (CAMERA 1, MINUS…);
  * meio ("via")    = texto solto logo acima de uma seta (PR RX SP-1…);
  * sentido         = ponta da seta na mesma altura do sinal;
  * VIDEO / AUDIO   = rótulos que dividem o slide em duas metades;
  * canais          = caixas com linhas "CH 1 – …";
  * "GRAVAR"        = marca o sinal logo acima como saída de gravação.

As notas do slide (campo "Clique para adicionar anotações" no PowerPoint)
permitem ajustar o resultado sem mexer no código, uma informação por linha:

    Nome: SP · Estúdio 6 SP com controle ION
    Apelidos: câmeras sp, narração sp
    Praça: SP
    ID: ion-sp-6sp
    Ocultar: sim

Use `python sincronizar.py --inspecionar arquivo.pptx` para ver como um
arquivo está sendo lido.
"""

from __future__ import annotations

import re
import statistics
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

# --------------------------------------------------------------------------
# Leitura genérica das formas do slide
# --------------------------------------------------------------------------


@dataclass
class Forma:
    id: int
    nome: str
    x: float
    y: float
    w: float
    h: float
    texto: str = ""
    linhas: list[str] = field(default_factory=list)
    preenchida: bool = False
    riscado: bool = False
    vermelho: bool = False
    titulo: bool = False
    # Para linhas/setas: pontos inicial e final e onde há ponta de seta.
    seta: tuple[float, float, float, float] | None = None
    ponta_inicio: bool = False
    ponta_fim: bool = False

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def x2(self) -> float:
        return self.x + self.w

    @property
    def y2(self) -> float:
        return self.y + self.h

    def resumo(self) -> dict:
        d = {
            "id": self.id,
            "nome": self.nome,
            "x": round(self.x),
            "y": round(self.y),
            "w": round(self.w),
            "h": round(self.h),
        }
        if self.texto:
            d["texto"] = self.texto
        for k in ("preenchida", "riscado", "vermelho", "titulo"):
            if getattr(self, k):
                d[k] = True
        if self.seta:
            d["seta"] = [round(v) for v in self.seta]
            d["pontas"] = ("◀" if self.ponta_inicio else "─") + ("▶" if self.ponta_fim else "─")
        return d


EMU_POR_PT = 12700
_TIPOS_LINHA = {"line", "straightConnector1", "bentConnector2", "bentConnector3", "curvedConnector3"}


def _normalizar(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFD", texto)
    sem_acento = "".join(c for c in sem_acento if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", sem_acento).strip().upper()


def _limpar(texto: str) -> str:
    return re.sub(r"\s+", " ", texto.replace("\u000b", " ")).strip()


def _xfrm_filhos(grupo) -> tuple[float, float, float, float, float, float]:
    """Transformação (escala e deslocamento) que o grupo aplica nas formas filhas."""
    xfrm = grupo._element.grpSpPr.find(qn("a:xfrm"))
    if xfrm is None:
        return 1, 1, 0, 0, 0, 0
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if None in (off, ext, ch_off, ch_ext):
        return 1, 1, 0, 0, 0, 0
    ccx, ccy = int(ch_ext.get("cx")), int(ch_ext.get("cy"))
    sx = int(ext.get("cx")) / ccx if ccx else 1
    sy = int(ext.get("cy")) / ccy if ccy else 1
    return sx, sy, int(off.get("x")), int(off.get("y")), int(ch_off.get("x")), int(ch_off.get("y"))


def _cor_vermelha(font) -> bool:
    try:
        if font.color is None or font.color.type is None:
            return False
        rgb = font.color.rgb
    except (AttributeError, TypeError):
        return False
    if rgb is None:
        return False
    r, g, b = rgb[0], rgb[1], rgb[2]
    return r >= 0xB0 and g <= 0x70 and b <= 0x70


def _ler_texto(sh, forma: Forma) -> None:
    if not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
        return
    linhas, runs = [], []
    for p in sh.text_frame.paragraphs:
        # quebras de linha (Shift+Enter) dentro do parágrafo também separam linhas
        for parte in re.split(r"[\u000b\n]", "".join(r.text for r in p.runs)):
            if parte.strip():
                linhas.append(_limpar(parte))
        runs.extend(r for r in p.runs if r.text.strip())
    forma.linhas = linhas
    forma.texto = " ".join(linhas)
    if runs:
        forma.riscado = all(
            (r.font._rPr is not None and r.font._rPr.get("strike") in ("sngStrike", "dblStrike")) for r in runs
        )
        forma.vermelho = any(_cor_vermelha(r.font) for r in runs)


def _ler_seta(sh, forma: Forma, el) -> None:
    geom = el.find(".//" + qn("a:prstGeom"))
    eh_conector = el.tag == qn("p:cxnSp")
    if not eh_conector and (geom is None or geom.get("prst") not in _TIPOS_LINHA):
        return
    xfrm = el.find(".//" + qn("a:xfrm"))
    flip_h = xfrm is not None and xfrm.get("flipH") == "1"
    flip_v = xfrm is not None and xfrm.get("flipV") == "1"
    x1, x2 = (forma.x2, forma.x) if flip_h else (forma.x, forma.x2)
    y1, y2 = (forma.y2, forma.y) if flip_v else (forma.y, forma.y2)
    forma.seta = (x1, y1, x2, y2)
    ln = el.find(".//" + qn("a:ln"))
    if ln is not None:
        head, tail = ln.find(qn("a:headEnd")), ln.find(qn("a:tailEnd"))
        # headEnd fica no ponto inicial da linha; tailEnd no final.
        forma.ponta_inicio = head is not None and head.get("type", "none") != "none"
        forma.ponta_fim = tail is not None and tail.get("type", "none") != "none"


def ler_formas(slide) -> list[Forma]:
    """Todas as formas do slide, com grupos desfeitos e coordenadas absolutas (EMU)."""
    formas: list[Forma] = []

    def visitar(shapes, transf):
        sx, sy, ox, oy, cx0, cy0 = transf
        for sh in shapes:
            if sh.left is None or sh.top is None:
                continue
            x = ox + (sh.left - cx0) * sx
            y = oy + (sh.top - cy0) * sy
            w = (sh.width or 0) * sx
            h = (sh.height or 0) * sy
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                gsx, gsy, gox, goy, gcx, gcy = _xfrm_filhos(sh)
                # compõe a transformação do grupo com a do grupo pai
                visitar(sh.shapes, (sx * gsx, sy * gsy, x - gcx * sx * gsx, y - gcy * sy * gsy, 0, 0))
                continue
            f = Forma(id=sh.shape_id, nome=sh.name, x=x, y=y, w=w, h=h)
            _ler_texto(sh, f)
            try:
                f.preenchida = sh.fill.type in (MSO_FILL.SOLID, MSO_FILL.GRADIENT, MSO_FILL.PATTERNED)
            except (AttributeError, TypeError, NotImplementedError):
                f.preenchida = False
            if not f.preenchida:
                # preenchimento herdado do estilo do tema (forma padrão do PowerPoint)
                style = sh._element.find(qn("p:style"))
                sp_pr = sh._element.find(qn("p:spPr"))
                sem_fill_explicito = sp_pr is None or sp_pr.find(qn("a:noFill")) is None
                f.preenchida = bool(
                    style is not None
                    and style.find(qn("a:fillRef")) is not None
                    and style.find(qn("a:fillRef")).get("idx", "0") != "0"
                    and sem_fill_explicito
                    and sp_pr is not None
                    and sp_pr.find(qn("a:solidFill")) is None
                    and sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX
                )
            if sh.is_placeholder:
                try:
                    f.titulo = sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
                except ValueError:
                    pass
            _ler_seta(sh, f, sh._element)
            formas.append(f)

    visitar(slide.shapes, (1, 1, 0, 0, 0, 0))
    return formas


def ler_notas(slide) -> dict[str, str]:
    if not slide.has_notes_slide:
        return {}
    texto = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
    campos = {}
    for linha in texto.splitlines():
        m = re.match(r"\s*([^:]{2,20}):\s*(.+)$", linha)
        if m:
            campos[_normalizar(m.group(1)).lower()] = m.group(2).strip()
    return campos


# --------------------------------------------------------------------------
# Interpretação do diagrama
# --------------------------------------------------------------------------

RE_CANAL = re.compile(r"^CH\s*(\d+)\s*[-–—:]?\s*(.*)$", re.I)
SECOES = {"VIDEO": "video", "AUDIO": "audio"}


class SlideIgnorado(Exception):
    """O slide não segue o padrão de modelo de evento (capa, logo…)."""


def _slug(texto: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", _normalizar(texto).lower()).strip("-")


def _eh_gravar(f: Forma) -> bool:
    return _normalizar(f.texto).replace(" ", "") == "GRAVAR"


def _praca_do_titulo(titulo: str) -> str:
    t = _normalizar(titulo)
    m = re.search(r"ION\s*X\s*(.+?)(?:\(|$)", t)
    if not m:
        return ""
    palavras = m.group(1).split()
    if palavras and palavras[0] == "SITE" and len(palavras) > 1:
        return palavras[1]
    return palavras[0] if palavras else ""


def _ler_modelo(f: Forma) -> dict[str, str]:
    """Caixa "Modelo: / Estúdio: / Controle:" → dicionário."""
    out: dict[str, str] = {}
    ultima = None
    for linha in f.linhas:
        n = _normalizar(linha)
        m = re.match(r"^(MODELO|ESTUDIO|CONTROLE)\s*:\s*(.*)$", n)
        if m:
            chave = {"MODELO": "modelo", "ESTUDIO": "estudio", "CONTROLE": "controle"}[m.group(1)]
            valor = linha.split(":", 1)[1].strip()
            out[chave] = valor
            ultima = chave
        elif n.startswith("CONTROLE"):  # "Controles externos"
            out["controle"] = linha
            ultima = "controle"
        elif ultima:  # continuação de uma linha quebrada
            out[ultima] = f"{out[ultima]} {linha}".strip()
    return {k: v for k, v in out.items() if v}


def _canais(f: Forma) -> dict | None:
    linhas_ch = [RE_CANAL.match(_limpar(l)) for l in f.linhas]
    if not any(linhas_ch):
        return None
    titulo = next((l for l, m in zip(f.linhas, linhas_ch) if not m), "").rstrip(":").strip()
    canais: dict[int, str] = {}
    for m in linhas_ch:
        if m:
            canais[int(m.group(1))] = m.group(2).strip()
    n = max(canais) if canais else 0
    grupo = {"titulo": titulo or "Canais", "canais": [canais.get(i, "") for i in range(1, n + 1)]}
    if f.vermelho:
        grupo["destaque"] = True
    return grupo


def _sinal(f: Forma | None, **extra) -> dict:
    if f is None or not f.texto:
        return {"nome": "", "vazio": True}
    s = {"nome": f.texto}
    if f.riscado:
        s["riscado"] = True
    if f.vermelho:
        s["destaque"] = True
    s.update({k: v for k, v in extra.items() if v})
    return s


def _sentido_da_seta(seta: Forma) -> int:
    """+1 se aponta para a direita, -1 para a esquerda, 0 se indefinido."""
    x1, _, x2, _ = seta.seta
    if seta.ponta_fim and not seta.ponta_inicio:
        return 1 if x2 > x1 else -1
    if seta.ponta_inicio and not seta.ponta_fim:
        return 1 if x1 > x2 else -1
    return 0


def _montar_fluxo(nos, pilulas, vias, setas, gravar, avisos, secao) -> dict:
    nos = sorted(nos, key=lambda n: n.cx)
    if len(nos) < 2:
        avisos.append(f"{secao}: encontrado(s) {len(nos)} bloco(s); são necessários pelo menos 2")
        return {"nos": [_sinal(n) for n in nos], "links": []}
    alt_pilula = statistics.median([p.h for p in pilulas]) if pilulas else 0
    tol = max(alt_pilula, 1)

    # "GRAVAR" marca só o sinal imediatamente acima dele
    gravar_ids = set()
    for g in gravar:
        acima = [p for p in pilulas if p.y2 <= g.y + tol / 2 and p.x < g.x2 and p.x2 > g.x]
        if acima:
            gravar_ids.add(max(acima, key=lambda p: p.y2).id)

    links = []
    for a, b in zip(nos, nos[1:]):
        no_vao = [p for p in pilulas if a.cx < p.cx < b.cx]
        vias_vao = [v for v in vias if a.x2 - tol < v.cx < b.x + tol]
        setas_vao = [s for s in setas if min(s.seta[0], s.seta[2]) < b.x and max(s.seta[0], s.seta[2]) > a.x2 - tol]
        ida, volta = [], []
        usadas: set[int] = set()
        for p in sorted(no_vao, key=lambda p: p.cy):
            # seta horizontal mais próxima, na altura do sinal
            candidatas = [
                s
                for s in setas_vao
                if abs((s.seta[1] + s.seta[3]) / 2 - p.cy) <= tol and abs(s.seta[1] - s.seta[3]) < tol
            ]
            seta = min(candidatas, key=lambda s: abs((s.seta[1] + s.seta[3]) / 2 - p.cy), default=None)
            sentido = _sentido_da_seta(seta) if seta else 0
            if not sentido:  # sem seta: sinal perto do bloco da esquerda sai dele
                sentido = 1 if p.cx - a.x2 < b.x - p.cx else -1
            # meio de transporte: texto logo acima da seta, do lado oposto ao sinal
            via = None
            y_ref = (seta.seta[1] + seta.seta[3]) / 2 if seta else p.cy
            cand_via = [
                v
                for v in vias_vao
                if v.id not in usadas
                and -tol * 1.6 <= v.cy - y_ref <= tol * 0.5
                and ((v.cx > p.cx) if sentido > 0 else (v.cx < p.cx))
            ]
            if cand_via:
                via = min(cand_via, key=lambda v: (abs(v.cy - y_ref), abs(v.cx - p.cx)))
                usadas.add(via.id)
            sinal = _sinal(p, via=via.texto if via else "", gravar=p.id in gravar_ids)
            (ida if sentido > 0 else volta).append(sinal)
        links.append({"ida": ida, "volta": volta})
    blocos = [{"nome": "GRAVAR", "gravar": True} if _eh_gravar(n) else _sinal(n) for n in nos]
    return {"nos": blocos, "links": links}


def interpretar_slide(slide, numero: int, arquivo: str) -> tuple[dict, list[str]]:
    """Converte um slide em evento. Levanta SlideIgnorado se não for um modelo."""
    formas = ler_formas(slide)
    notas = ler_notas(slide)
    avisos: list[str] = []
    if _normalizar(notas.get("ocultar", "")) in ("SIM", "S", "TRUE", "1", "X"):
        raise SlideIgnorado("marcado com 'Ocultar: sim' nas notas")

    rotulos = {SECOES[_normalizar(f.texto)]: f for f in formas if _normalizar(f.texto) in SECOES}
    if not rotulos:
        raise SlideIgnorado("sem os rótulos VIDEO/AUDIO")

    restantes = [f for f in formas if f not in rotulos.values()]

    # Título: placeholder de título ou o texto mais alto do slide.
    com_texto = [f for f in restantes if f.texto]
    titulo_f = next((f for f in com_texto if f.titulo), None) or min(
        (f for f in com_texto if f.y2 < min(r.y for r in rotulos.values())), key=lambda f: f.y, default=None
    )
    titulo = titulo_f.texto if titulo_f else f"Slide {numero}"
    restantes = [f for f in restantes if f is not titulo_f]

    evento: dict = {"titulo": titulo}
    canais, notas_slide = [], []
    gravar, livres = [], []
    for f in restantes:
        n = _normalizar(f.texto)
        if not f.texto:
            if f.seta or f.preenchida:
                livres.append(f)
            continue
        if re.search(r"\b(ESTUDIO|CONTROLE)S?\b", n) and ":" in f.texto and n.startswith(("MODELO", "ESTUDIO")):
            evento.update(_ler_modelo(f))
        elif (grupo := _canais(f)) is not None:
            canais.append((f, grupo))
        elif "EM CONSTRU" in n:
            evento["emConstrucao"] = True
        elif n.startswith("STATION"):
            evento["station"] = f.texto.split(":", 1)[-1].strip()
        elif f.texto.startswith("*"):
            notas_slide.append(f.texto)
        elif _eh_gravar(f) and f.h < (max((r.h for r in rotulos.values()), default=0) * 3):
            gravar.append(f)
        else:
            livres.append(f)

    # Divide o slide em VÍDEO (em cima) e ÁUDIO (embaixo).
    if "video" in rotulos and "audio" in rotulos:
        corte = rotulos["audio"].y
    elif "audio" in rotulos:
        corte = -1
    else:
        corte = float("inf")

    setas = [f for f in livres if f.seta]
    caixas = [f for f in livres if not f.seta]
    for secao in ("video", "audio"):
        if secao not in rotulos:
            continue
        na_secao = [f for f in caixas if (f.cy < corte) == (secao == "video")]
        setas_secao = [f for f in setas if ((f.seta[1] + f.seta[3]) / 2 < corte) == (secao == "video")]
        gravar_secao = [g for g in gravar if (g.cy < corte) == (secao == "video")]
        if not na_secao:
            continue
        # Bloco = caixa bem mais alta que o sinal típico.
        alturas = sorted(f.h for f in na_secao if f.preenchida) or sorted(f.h for f in na_secao)
        alt_tipica = statistics.median(alturas)
        nos = [f for f in na_secao if f.h > alt_tipica * 2.5 and f.texto]
        resto = [f for f in na_secao if f not in nos]
        pilulas = [f for f in resto if f.preenchida]
        vias = [f for f in resto if not f.preenchida and f.texto]
        evento[secao] = _montar_fluxo(nos, pilulas, vias, setas_secao, gravar_secao, avisos, secao)

    if canais:
        # ordem de leitura: colunas da esquerda para a direita, de cima para baixo
        canais.sort(key=lambda fg: (round(fg[0].x / (fg[0].w or 1)), fg[0].y))
        evento["canais"] = [g for _, g in canais]
    if notas_slide:
        evento["notas"] = notas_slide

    # Campos derivados e ajustes vindos das notas do slide.
    evento["praca"] = (notas.get("praca") or _praca_do_titulo(titulo)).upper()
    if "nome" in notas:
        evento["nome"] = notas["nome"]
    else:
        partes = [p for p in (evento.get("estudio") and f"Estúdio {evento['estudio']}",
                              evento.get("controle") and f"controle {evento['controle']}") if p]
        sub = re.search(r"\(([^)]+)\)", titulo)
        base = f"{evento['praca']} · " if evento["praca"] else ""
        evento["nome"] = base + (f"{sub.group(1)} — " if sub else "") + (" com ".join(partes) or titulo)
    if "apelidos" in notas:
        evento["apelidos"] = [a.strip() for a in notas["apelidos"].split(",") if a.strip()]
    evento["id"] = notas.get("id") or f"{_slug(Path(arquivo).stem)}-s{slide.slide_id}"
    evento["fonte"] = {"arquivo": arquivo, "slide": numero, "slide_id": slide.slide_id}

    for secao in ("video", "audio"):
        if secao in evento and not any(l["ida"] or l["volta"] for l in evento[secao]["links"]):
            avisos.append(f"{secao}: nenhum sinal encontrado entre os blocos")
    return evento, avisos


def _anotacoes(slide) -> list[str] | None:
    formas = ler_formas(slide)
    textos = sorted((f for f in formas if f.texto), key=lambda f: (f.y, f.x))
    if not textos or not _normalizar(textos[0].texto).startswith("ANOTA"):
        return None
    return [l for f in textos[1:] for l in f.linhas]


def extrair(caminho: Path, arquivo_rel: str) -> dict:
    """
    Lê um .pptx e devolve {"eventos", "anotacoes", "avisos", "ignorados"}.
    "ignorados" lista slides que não são modelos (capa, logo…) — só informativo.
    `arquivo_rel` é o caminho relativo à pasta sincronizada (vai para o JSON).
    """
    prs = Presentation(str(caminho))
    eventos, anotacoes, avisos, ignorados = [], [], [], []
    for numero, slide in enumerate(prs.slides, start=1):
        try:
            ev, av = interpretar_slide(slide, numero, arquivo_rel)
            eventos.append(ev)
            avisos.extend(f"{arquivo_rel} slide {numero}: {a}" for a in av)
        except SlideIgnorado as e:
            notas = _anotacoes(slide)
            if notas:
                anotacoes.extend(notas)
            else:
                ignorados.append(f"slide {numero}: {e}")
    return {"eventos": eventos, "anotacoes": anotacoes, "avisos": avisos, "ignorados": ignorados}


def inspecionar(caminho: Path) -> dict:
    """Despejo bruto das formas de cada slide — para calibrar a extração."""
    prs = Presentation(str(caminho))
    return {
        "largura_emu": prs.slide_width,
        "altura_emu": prs.slide_height,
        "slides": [
            {
                "numero": i,
                "slide_id": s.slide_id,
                "layout": s.slide_layout.name,
                "notas": ler_notas(s),
                "formas": [f.resumo() for f in ler_formas(s)],
            }
            for i, s in enumerate(prs.slides, start=1)
        ],
    }
